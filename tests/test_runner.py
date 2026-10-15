from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from core.runner import StageError, run_stage, save_result
from core.settings import SEVERITY_FILLS, WARNING
from tests import stage_99_fake
from tests.helpers import make_xlsx

HEADER = ["Codigo", "Fecha", "Valor", "Fuente"]


def area_a():
    return make_xlsx(
        [
            HEADER,
            ["001", "05/03/2024", "1.500.000", " PGN "],
            [None, None, None, None],
            ["002", datetime(2024, 3, 6), 2000000, "Otros"],
        ],
        sheet="Ejecucion",
        name="area_a.xlsx",
    )


def area_b():
    return make_xlsx(
        [HEADER, ["001", "05/03/2024", "1500000", "PGN"], ["003", "xx/03/2024", "abc", "PGN"]],
        sheet="Ejecucion",
        name="area_b.xlsx",
    )


def catalog():
    return make_xlsx([["Codigo", "Nombre"], ["001", "Proyecto uno"]], sheet="Catalogo", header_row=2, name="cat.xlsx")


def test_full_run():
    lines = []
    result = run_stage(
        stage_99_fake,
        {"reports": [area_a(), area_b()], "catalog": [catalog()]},
        progress=lines.append,
    )
    summary = result.summary
    # El log de ejecución se recibe en vivo, línea por línea, igual al que queda guardado.
    assert lines == result.audit.step_lines()
    titles = [step.title for step in result.audit.steps]
    assert titles[:4] == [
        "Verificando archivos cargados", "Abriendo «area_a.xlsx»",
        "Validando la estructura de «area_a.xlsx»", "Abriendo «area_b.xlsx»",
    ]
    assert "Eliminando las filas de título del reporte" in titles
    assert titles[-1] == "Generando el archivo de salida y el log"
    assert all(step.details for step in result.audit.steps)
    assert "   ✓ 3 archivos: area_a.xlsx, area_b.xlsx, cat.xlsx" in lines
    assert "   ✓ 4 columnas, 3 registros" in lines
    # Con varios archivos, las filas se indican junto con el nombre del archivo.
    titles_step = result.audit.steps[titles.index("Eliminando las filas de título del reporte")]
    assert titles_step.details == ["✓ 1 fila eliminada: cat.xlsx: fila 1"]
    duplicates = result.audit.steps[titles.index("Eliminando duplicados")]
    assert duplicates.details == ["✓ 1 fila eliminada: area_b.xlsx: fila 2"]
    assert summary.rows_read == 6  # 3 + 2 de los reportes y 1 del catálogo
    assert summary.rows_removed == 2  # Una fila vacía y un duplicado entre archivos
    assert summary.rows_output == 3
    assert [(i.file_name, i.rows) for i in summary.input_files] == [
        ("area_a.xlsx", 3), ("area_b.xlsx", 2), ("cat.xlsx", 1)
    ]
    assert result.output_file_name.startswith("E99_fake_procesado_")
    assert result.log_file_name.startswith("E99_fake_log_")

    # Archivo de salida: sin columnas internas, con formatos y resaltado.
    sheet = load_workbook(BytesIO(result.output_bytes)).active
    header = [c.value for c in sheet[1]]
    assert header == ["Codigo", "Fecha", "Valor", "Fuente", "Nombre", "Valor (millones)"]
    assert sheet["A2"].value == "001" and sheet["E2"].value == "Proyecto uno"
    assert sheet["B2"].number_format == "dd/mm/yyyy"
    assert sheet["C2"].number_format == "#,##0.00"
    assert sheet.auto_filter.ref and sheet.freeze_panes == "A2"
    assert sheet["D3"].fill.start_color.rgb.endswith(SEVERITY_FILLS[WARNING])  # Fuente "Otros"
    assert sheet["C4"].value == "abc" and sheet["C4"].fill.fill_type == "solid"  # Valor inválido

    # Log: seis hojas con filtros y paneles inmovilizados.
    log = load_workbook(BytesIO(result.log_bytes))
    assert log.sheetnames == ["Resumen", "Ejecución", "Acciones", "Hallazgos", "Cambios", "Eliminados"]
    execution = list(log["Ejecución"].iter_rows(min_row=2, values_only=True))
    assert execution[0] == ("Verificando archivos cargados", "✓ 3 archivos: area_a.xlsx, area_b.xlsx, cat.xlsx")
    assert execution[-1][1].startswith("✓ Terminado en ")
    for ws in log.worksheets:
        assert ws.auto_filter.ref and ws.freeze_panes == "A2"
    removed_rows = [(r[0], r[1]) for r in log["Eliminados"].iter_rows(min_row=2, values_only=True)]
    # Primero la fila de título del catálogo (encabezados en la fila 2), luego los registros.
    assert removed_rows == [("cat.xlsx", 1), ("area_a.xlsx", 3), ("area_b.xlsx", 2)]
    assert summary.title_rows_removed == 1


def test_missing_required_input():
    with pytest.raises(StageError) as error:
        run_stage(stage_99_fake, {"catalog": [catalog()]})
    assert any("Reportes de prueba" in p for p in error.value.problems)


def test_too_many_files_for_single_input():
    with pytest.raises(StageError) as error:
        run_stage(stage_99_fake, {"reports": [area_a()], "catalog": [catalog(), catalog()]})
    assert any("un solo archivo" in p for p in error.value.problems)


def test_structure_errors_from_all_files_stop_processing():
    bad = make_xlsx([["Codigo", "Valor"]], sheet="Ejecucion", name="malo.xlsx")
    wrong_sheet = make_xlsx([HEADER], sheet="Hoja1", name="otra_hoja.xlsx")
    with pytest.raises(StageError) as error:
        run_stage(stage_99_fake, {"reports": [area_a(), bad, wrong_sheet]})
    problems = error.value.problems
    assert error.value.message == "La estructura de 2 de los 3 archivos no es la esperada."
    assert any("malo.xlsx" in p and "EST-003" in p for p in problems)
    assert any("otra_hoja.xlsx" in p and "EST-002" in p for p in problems)
    # El log de ejecución muestra hasta dónde llegó y en qué paso falló.
    steps = error.value.steps
    assert "▸ Validando la estructura de «malo.xlsx»" in steps
    assert any(line.startswith("   ✗ [EST-003] malo.xlsx") for line in steps)
    assert "▸ Generando el archivo de salida y el log" not in steps


def test_file_that_is_not_xlsx_fails_when_opening():
    not_excel = BytesIO(b"esto no es un excel")
    not_excel.name = "texto.xlsx"
    with pytest.raises(StageError) as error:
        run_stage(stage_99_fake, {"reports": [not_excel]})
    steps = error.value.steps
    assert steps[-2:] == [
        "▸ Abriendo «texto.xlsx»",
        "   ✗ [EST-001] texto.xlsx: no se pudo abrir. Verifique que sea un archivo .xlsx válido.",
    ]


def test_missing_input_is_shown_in_step_log():
    with pytest.raises(StageError) as error:
        run_stage(stage_99_fake, {"catalog": [catalog()]})
    assert error.value.steps[0] == "▸ Verificando archivos cargados"
    assert error.value.steps[1].startswith("   ✗ [EST-007]")


def test_save_result_never_overwrites(tmp_path):
    result = run_stage(stage_99_fake, {"reports": [area_a()]})
    first = save_result(result, tmp_path)
    second = save_result(result, tmp_path)
    assert first[0] != second[0] and first[1] != second[1]
    assert second[0].name.endswith("_1.xlsx")
    assert len(list(tmp_path.iterdir())) == 4
