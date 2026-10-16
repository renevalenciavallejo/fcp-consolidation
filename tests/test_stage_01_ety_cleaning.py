"""Etapa 1: validación de estructura del reporte ETY, con datos inventados."""

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from core.runner import StageError, run_stage
from stages import stage_01_ety_cleaning
from tests.helpers import make_xlsx

COLUMNS = list(stage_01_ety_cleaning.INPUTS["ety_report"]["columns"])
# Columnas del archivo de salida: las del reporte más Diferencia a la derecha de Importe (K).
OUTPUT_COLUMNS = [*COLUMNS[:10], "Diferencia", *COLUMNS[10:]]

ROW = [
    "0010000001", "12345678", "PERSONA DE PRUEBA", datetime(2026, 9, 28, 2, 0), 1000000, 0, 0,
    5000, 0, 995000, "AP00000001", "Cuentas a Pagar", 1, "00000001", "Pago Estándar",
    "CE0000000001", "DS_PRUEBA", "20025", "23-26", "1002", "4147", "2", "100212", "0002",
]


TITLE_ROWS = [
    ["Todos Pagos X Comprobante", " 1"],
    ["De Fecha = 2026-09-28"],
    ["A Fecha = 2026-09-30"],
]


def distinct(*rows):
    """Copia las filas con un Comprobante distinto cada una, para que no se depuren como duplicados."""
    result = []
    for number, row in enumerate(rows, start=101):
        row = list(row)
        row[COLUMNS.index("Comprobante")] = f"{number:08d}"
        result.append(row)
    return tuple(result)


def ety_file(header=COLUMNS, rows=(ROW,), sheet="sheet1", titles=TITLE_ROWS, name="ety.xlsx"):
    return make_xlsx([*titles, list(header), *[list(r) for r in rows]], sheet=sheet, name=name)


def run(*files):
    return run_stage(stage_01_ety_cleaning, {"ety_report": list(files)})


def stage_error(*files) -> StageError:
    with pytest.raises(StageError) as error:
        run(*files)
    return error.value


def problems_for(*files) -> list[str]:
    return stage_error(*files).problems


def test_messages_speak_of_a_single_file():
    error = stage_error(ety_file(titles=[]))
    assert error.message == "La estructura del archivo no es la esperada."


def test_valid_report_is_processed():
    result = run(ety_file(rows=distinct(ROW, ROW)))
    assert result.summary.rows_read == 2
    assert result.output_file_name == "ety_E01_procesado.xlsx"
    assert result.log_file_name == "ety_E01_log.xlsx"


def test_header_spaces_are_tolerated():
    header = [f" {c} " for c in COLUMNS]
    assert run(ety_file(header=header)).summary.rows_read == 1


def test_any_sheet_name_is_accepted():
    assert run(ety_file(sheet="Hoja1")).summary.rows_read == 1


def test_output_keeps_the_original_sheet_name():
    result = run(ety_file(sheet="sheet1"))
    assert load_workbook(BytesIO(result.output_bytes)).sheetnames == ["sheet1"]


def test_column_order_is_validated():
    header = list(COLUMNS)
    header[4], header[5] = header[5], header[4]  # VR. BRUTO <-> VR. RTEFTE
    problems = problems_for(ety_file(header=header))
    assert len(problems) == 1
    assert "EST-008" in problems[0] and "posición 5" in problems[0] and "'VR. BRUTO'" in problems[0]


def test_missing_column():
    header = [c for c in COLUMNS if c != "Importe"]
    row = [v for c, v in zip(COLUMNS, ROW) if c != "Importe"]
    problems = problems_for(ety_file(header=header, rows=(row,)))
    assert any("EST-003" in p and "'Importe'" in p for p in problems)


def test_extra_column():
    problems = problems_for(ety_file(header=[*COLUMNS, "Otra"], rows=([*ROW, "x"],)))
    assert any("EST-004" in p and "'Otra'" in p for p in problems)


def test_title_rows_removed():
    problems = problems_for(ety_file(titles=[]))
    assert len(problems) == 1 and "están en la fila 1" in problems[0]


def test_only_one_file_allowed():
    problems = problems_for(ety_file(name="a.xlsx"), ety_file(name="b.xlsx"))
    assert any("EST-007" in p and "un solo archivo" in p for p in problems)


def test_title_rows_are_removed_and_logged():
    result = run(ety_file(rows=distinct(ROW, ROW)))
    summary = result.summary
    assert summary.rows_output == 2
    assert summary.rows_removed == 0          # Las filas de título no son registros
    assert summary.title_rows_removed == 3

    removals = result.audit.removals_df()
    assert list(removals["Fila origen"]) == [1, 2, 3]
    assert set(removals["Código de regla"]) == {"LIM-003"}
    assert removals.iloc[0]["Beneficiario Pago"] == "Todos Pagos X Comprobante"  # Columna A
    assert removals.iloc[0]["Identificación"] == " 1"                            # Columna B

    sheet = load_workbook(BytesIO(result.output_bytes)).active
    assert [c.value for c in sheet[1]] == OUTPUT_COLUMNS  # La salida empieza con los encabezados
    assert sheet["A2"].value == ROW[0]


CURRENCY_COLUMNS = ["VR. BRUTO", "VR. RTEFTE", "VR. RETEIVA", "VR. RETEICA", "VR. TASAS", "Importe"]


def test_columns_e_to_j_are_currency():
    assert COLUMNS[4:10] == CURRENCY_COLUMNS  # Columnas E a J del reporte


def test_currency_columns_are_converted_and_formatted():
    row = list(ROW)
    row[COLUMNS.index("VR. BRUTO")] = "$ 1.000.000"   # Texto con formato colombiano
    row[COLUMNS.index("Importe")] = "995000"          # Texto con un número simple
    result = run(ety_file(rows=distinct(ROW, row)))

    sheet = load_workbook(BytesIO(result.output_bytes)).active
    for letter in "EFGHIJ":
        for number in (2, 3):
            cell = sheet[f"{letter}{number}"]
            assert isinstance(cell.value, (int, float)), cell.coordinate  # Excel guarda 1000000.0 como 1000000
            assert cell.number_format == '"$" #,##0.00', cell.coordinate
    assert sheet["E3"].value == 1000000.0
    assert sheet["D2"].number_format != '"$" #,##0.00'  # La fecha no es moneda

    # Solo se registra el cambio cuyo valor visible cambió.
    changes = [(c.source_row, c.column, c.old_value, c.new_value) for c in result.audit.changes]
    assert changes == [(6, "VR. BRUTO", "$ 1.000.000", 1000000.0)]
    assert result.audit.changes[0].rule == "TIP-005"


def test_invalid_currency_value_is_reported():
    row = list(ROW)
    row[COLUMNS.index("VR. RTEFTE")] = "sin valor"
    result = run(ety_file(rows=(row,)))

    findings = [(f.source_row, f.column, f.severity, f.rule) for f in result.audit.findings]
    assert findings == [(5, "VR. RTEFTE", "Error", "TIP-005")]
    sheet = load_workbook(BytesIO(result.output_bytes)).active
    assert sheet["F2"].value == "sin valor"           # Se conserva el valor original
    assert sheet["F2"].fill.fill_type == "solid"      # Resaltado como error



def with_description(description):
    row = list(ROW)
    row[COLUMNS.index("Descripción")] = description
    return row


def test_cash_distribution_rows_are_removed_and_logged():
    rows = distinct(
        with_description("Cuentas a Pagar"),
        with_description("Distribución Caja"),
        with_description(" distribucion  CAJA "),   # Variación de escritura
        with_description("Cuentas a Pagar"),
    )
    result = run(ety_file(rows=rows))

    assert result.summary.rows_read == 4
    assert result.summary.rows_output == 2
    assert result.summary.rows_removed == 2

    removed = [r for r in result.audit.removals if r.rule == "LIM-101"]
    assert [r.source_row for r in removed] == [6, 7]
    assert removed[0].reason == "Descripción «Distribución Caja»"
    assert removed[0].row_data["Beneficiario Pago"] == ROW[0]  # Se guarda el registro completo

    sheet = load_workbook(BytesIO(result.output_bytes)).active
    descriptions = [sheet[f"M{r}"].value for r in range(2, sheet.max_row + 1)]  # L se corrió a M
    assert descriptions == ["Cuentas a Pagar", "Cuentas a Pagar"]



def test_difference_column_has_formula_right_of_importe():
    rows = distinct(with_description("Distribución Caja"), ROW, ROW)
    result = run(ety_file(rows=rows))

    sheet = load_workbook(BytesIO(result.output_bytes)).active
    assert sheet["J1"].value == "Importe" and sheet["K1"].value == "Diferencia"
    assert sheet["L1"].value == "Asiento"
    # La fila de Distribución Caja se eliminó: las fórmulas usan la fila de salida.
    assert sheet["K2"].value == "=E2-F2-G2-H2-I2-J2"
    assert sheet["K3"].value == "=E3-F3-G3-H3-I3-J3"
    assert sheet["K2"].number_format == '"$" #,##0.00'

    added = result.audit.added_columns
    assert [(a.column, a.rule, a.rows) for a in added] == [("Diferencia", "TRA-101", 2)]
    assert result.summary.cells_changed == 0  # La columna nueva no cuenta como celdas modificadas

    summary = load_workbook(BytesIO(result.log_bytes))["Resumen"]
    sections = [row for row in summary.iter_rows(values_only=True) if row[0] == "Columnas agregadas"]
    assert len(sections) == 1 and sections[0][1].startswith("Diferencia (TRA-101") and sections[0][2] == "2 filas"



def with_voucher(voucher, description="Cuentas a Pagar", **changes):
    """Fila con el Comprobante dado; `changes` cambia otras columnas (por ejemplo Cuenta="9999")."""
    row = with_description(description)
    row[COLUMNS.index("Comprobante")] = voucher
    for column, value in changes.items():
        row[COLUMNS.index(column.replace("_", " "))] = value
    return row


def test_repeated_vouchers_are_flagged_and_rows_sorted():
    rows = (
        with_voucher("00000003"),                       # fila 5
        with_voucher("00000001"),                       # fila 6
        with_voucher("00000003", Cuenta="9999"),        # fila 7 (repetido, distinta Cuenta: se conserva)
        with_voucher("00000002"),                       # fila 8
        with_voucher("00000002", "Distribución Caja"),  # fila 9: se elimina antes, no cuenta
    )
    result = run(ety_file(rows=rows))

    findings = [(f.source_row, f.column, f.severity, f.rule) for f in result.audit.findings]
    assert findings == [(5, "Comprobante", "Advertencia", "VAL-101"), (7, "Comprobante", "Advertencia", "VAL-101")]

    sheet = load_workbook(BytesIO(result.output_bytes)).active
    assert sheet["O1"].value == "Comprobante"
    assert [sheet[f"O{r}"].value for r in range(2, 6)] == ["00000001", "00000002", "00000003", "00000003"]
    # Los repetidos quedan juntos y resaltados en amarillo; la fórmula usa la fila nueva.
    assert sheet["O4"].fill.start_color.rgb.endswith("FFEB9C") and sheet["O5"].fill.start_color.rgb.endswith("FFEB9C")
    assert sheet["O2"].fill.fill_type is None
    assert sheet["K4"].value == "=E4-F4-G4-H4-I4-J4"

    transformations = [(t.rule, t.detail) for t in result.audit.transformations]
    assert transformations == [("TRA-102", "Registros ordenados por Comprobante (4 registros)")]


def test_step_log_shows_each_rule():
    cash = list(ROW)
    cash[11] = "Distribución Caja"
    lines = run(ety_file(rows=(ROW, cash, ROW))).audit.step_lines()  # Filas 5 y 7 idénticas
    assert lines[:4] == [
        "▸ Verificando archivos cargados",
        "   ✓ 1 archivo: ety.xlsx",
        "▸ Abriendo «ety.xlsx»",
        "   ✓ Archivo .xlsx válido",
    ]
    assert "   ✓ 24 columnas en el orden esperado, 3 registros" in lines
    assert "   ✓ 3 filas eliminadas: filas 1, 2 y 3" in lines
    titles = [line for line in lines if line.startswith("▸")]
    assert titles[4:] == [
        "▸ Convirtiendo a moneda las columnas VR. BRUTO a Importe",
        "▸ Eliminando las filas con Descripción «Distribución Caja»",
        "▸ Agregando la columna Diferencia",
        "▸ Ordenando por Comprobante",
        "▸ Eliminando registros duplicados por Comprobante (con Diferencia 0)",
        "▸ Revisando registros con Comprobante repetido",
        "▸ Generando el archivo de salida y el log",
    ]
    assert "   ✓ 1 fila eliminada: fila 6" in lines   # Distribución Caja
    assert "   ✓ 1 fila eliminada: fila 7" in lines   # Duplicado de la fila 5
    # Después de depurar, el comprobante ya no está repetido.
    assert lines[lines.index("▸ Revisando registros con Comprobante repetido") + 1] == "   ✓ Sin cambios"


def test_actions_table_rebuilds_what_was_done_in_order():
    cash = list(ROW)
    cash[11] = "Distribución Caja"
    actions = run(ety_file(rows=(ROW, cash, ROW))).audit.actions_df()
    assert actions["Acción"].tolist() == [
        "Fila eliminada", "Fila eliminada", "Fila eliminada",   # Filas de título 1 a 3
        "Fila eliminada",                                       # «Distribución Caja», fila 6
        "Columna agregada",                                     # Diferencia
        "Transformación",                                       # Orden por Comprobante
        "Fila eliminada",                                       # Duplicado de la fila 5 (fila 7)
    ]
    assert actions["Fila"].tolist()[:4] == [1, 2, 3, 6]
    assert actions.iloc[3]["Paso"] == "Eliminando las filas con Descripción «Distribución Caja»"
    # Cada acción muestra el registro completo, con todas las columnas del reporte.
    assert "Archivo" not in actions.columns
    assert all(column in actions.columns for column in COLUMNS)
    assert actions.iloc[3]["Descripción"] == "Distribución Caja"
    assert actions.iloc[6]["Fila"] == 7 and actions.iloc[6]["Regla"] == "DUP-101"



# --- Paso 7: eliminación de duplicados por Comprobante -------------------------


def dedup_run(*rows):
    result = run(ety_file(rows=rows))
    removed = [(r.source_row, r.rule) for r in result.audit.removals if r.rule.startswith("DUP")]
    return result, removed


def output_rows(result):
    sheet = load_workbook(BytesIO(result.output_bytes)).active
    columns = [c.value for c in sheet[1]]
    return [dict(zip(columns, [c.value for c in row])) for row in sheet.iter_rows(min_row=2)]


def test_case_2_1_identical_rows_keep_the_first():
    result, removed = dedup_run(with_voucher("1"), with_voucher("1"), with_voucher("1"))
    assert removed == [(6, "DUP-101"), (7, "DUP-101")]
    assert result.audit.removals[-1].reason == "Duplicado de la fila 5 (Comprobante 1, caso 2.1)"


def test_case_2_2_only_dpto_or_cd_fnd_differ():
    _, removed = dedup_run(with_voucher("1"), with_voucher("1", Dpto="99999", Cd_Fnd="99-99"))
    assert removed == [(6, "DUP-102")]


def test_case_2_3_blocking_columns_keep_all():
    for column in ("Subfideicomiso", "Cuenta", "Programa"):
        result, removed = dedup_run(with_voucher("1"), with_voucher("1", Dpto="99999", **{column: "X"}))
        assert removed == [], column  # 2.2 y 2.3 a la vez: prevalece 2.3
    notes = [line for line in result.audit.step_lines() if "caso 2.3" in line]
    assert notes == ["   ✓ Se conserva 1 grupo porque difieren en Subfideicomiso, Cuenta o Programa (caso 2.3): 1"]


def test_case_2_4_other_columns_keep_all():
    _, removed = dedup_run(with_voucher("1"), with_voucher("1", Proy="X"))
    assert removed == []


def test_only_rows_with_zero_difference_are_evaluated():
    unbalanced = with_voucher("1")
    unbalanced[COLUMNS.index("Importe")] = 990000  # Diferencia = 5000
    _, removed = dedup_run(with_voucher("1"), unbalanced, with_voucher("1"))
    assert removed == [(7, "DUP-101")]  # La fila 6 no participa y se conserva


def test_groups_are_not_required_to_be_contiguous_and_order_is_kept():
    result, removed = dedup_run(
        with_voucher("2", Proy="A"),   # fila 5
        with_voucher("1"),             # fila 6
        with_voucher("2", Proy="B"),   # fila 7: grupo 2 difiere en Proy (2.4), se conservan
        with_voucher("1"),             # fila 8: duplicado de la fila 6
    )
    assert removed == [(8, "DUP-101")]
    # Las filas que quedan conservan el orden que tenían (ya ordenado por Comprobante en el paso 6).
    assert [(r["Comprobante"], r["Proy"]) for r in output_rows(result)] == [("1", "2"), ("2", "A"), ("2", "B")]


def test_first_is_the_highest_row_in_the_original_file():
    # El orden por Comprobante no cambia cuál es «la primera»: la de menor fila de origen.
    _, removed = dedup_run(with_voucher("1", Dpto="B"), with_voucher("1", Dpto="A"))
    assert removed == [(6, "DUP-102")]


def test_kept_rows_of_a_cleaned_group_are_not_flagged_as_repeated():
    result, _ = dedup_run(with_voucher("1"), with_voucher("1"), with_voucher("2"), with_voucher("2", Cuenta="X"))
    flagged = [(f.source_row, f.rule) for f in result.audit.findings]
    assert flagged == [(7, "VAL-101"), (8, "VAL-101")]  # Solo el grupo 2, que sigue repetido


def test_difference_is_compared_to_the_cent():
    # Cuadra al centavo, pero la resta exacta deja un residuo de 0,0000000019 (caso real).
    values = {"VR. BRUTO": 10602576, "VR. RTEFTE": 486809.52, "VR. RETEIVA": 0,
              "VR. RETEICA": 77154.95, "VR. TASAS": 0, "Importe": 10038611.53}
    rows = []
    for _ in range(2):
        row = with_voucher("1")
        for column, value in values.items():
            row[COLUMNS.index(column)] = value
        rows.append(row)
    _, removed = dedup_run(*rows)
    assert removed == [(6, "DUP-101")]


def test_difference_of_one_cent_is_not_zero():
    unbalanced = with_voucher("1")
    unbalanced[COLUMNS.index("Importe")] = 994999.99  # Diferencia = 0,01
    _, removed = dedup_run(with_voucher("1"), unbalanced)
    assert removed == []


# Una segunda fila del mismo pago, como aparece en el reporte real: otra fecha, otro asiento,
# otra línea y otro número de documento.
IGNORED = {"FECHA DE PAGO": datetime(2026, 9, 29, 2, 0), "Asiento": "AP99999999", "LNP": 2,
           "Nº Sec Doc": "CE9999999999"}


def test_ignored_columns_are_not_compared_in_cases_2_1_and_2_2():
    # 2.1: solo difieren en FECHA DE PAGO, Asiento, LNP y Nº Sec Doc.
    result, removed = dedup_run(with_voucher("1"), with_voucher("1", **IGNORED))
    assert removed == [(6, "DUP-101")]
    kept = output_rows(result)[0]
    for column in IGNORED:  # La fila conservada mantiene sus valores originales
        assert kept[column] == ROW[COLUMNS.index(column)], column
    # Cada columna por separado tampoco cuenta.
    for column, value in IGNORED.items():
        _, removed = dedup_run(with_voucher("1"), with_voucher("1", **{column: value}))
        assert removed == [(6, "DUP-101")], column
    # 2.2: difieren en Dpto y en las columnas ignoradas.
    _, removed = dedup_run(with_voucher("1"), with_voucher("1", Dpto="99999", **IGNORED))
    assert removed == [(6, "DUP-102")]


def test_ignored_columns_do_not_override_case_2_3_or_2_4():
    _, removed = dedup_run(with_voucher("1"), with_voucher("1", Cuenta="X", **IGNORED))
    assert removed == []   # 2.3 prevalece
    _, removed = dedup_run(with_voucher("1"), with_voucher("1", Proy="X", **IGNORED))
    assert removed == []   # La diferencia en Proy sigue siendo 2.4
