from io import BytesIO

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from core.writer import ExcelFormula, build_output, capture_style


def test_formula_references_are_resolved_by_column_name():
    formula = ExcelFormula("={A}-{C}", "#,##0.00")
    df = pd.DataFrame({"A": [5.0, 7.0], "B": ["x", "y"], "C": [1.0, 2.0], "D": [formula, formula]}, dtype=object)
    sheet = load_workbook(BytesIO(build_output(df, []))).active
    assert sheet["D2"].value == "=A2-C2"
    assert sheet["D3"].value == "=A3-C3"
    assert sheet["D2"].number_format == "#,##0.00"


def test_formula_with_unknown_column_fails():
    df = pd.DataFrame({"A": [1.0], "D": [ExcelFormula("={Z}-{A}")]}, dtype=object)
    with pytest.raises(ValueError, match="'Z'"):
        build_output(df, [])


def styled_source():
    """Hoja con apariencia propia, como la de un reporte exportado por el sistema."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "sheet1"
    sheet.append(["Valor", "Cantidad"])
    sheet.append([1.5, 3])
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="FF969696")
        cell.border = Border(bottom=Side(style="double"))
    sheet["A2"].number_format = "0.000"
    sheet["B2"].number_format = "0"
    sheet["B2"].font = Font(name="Arial", size=9)
    sheet.column_dimensions["A"].width = 29
    return sheet


def test_original_appearance_is_kept():
    style = capture_style(styled_source(), 1, {"Valor": 0, "Cantidad": 1})
    df = pd.DataFrame({
        "Valor": [2.5], "Diferencia": [ExcelFormula("={Valor}", "#,##0.00")], "Cantidad": [4],
    }, dtype=object)
    sheet = load_workbook(BytesIO(build_output(df, [], {"Valor": '"$" #,##0.00'}, style))).active
    assert sheet.title == "sheet1"
    # Encabezados con el estilo original; la columna nueva toma el de su izquierda.
    for cell in sheet[1]:
        assert (cell.font.name, cell.font.b, cell.border.bottom.style) == ("Arial", True, "double"), cell.coordinate
        assert cell.fill.fgColor.rgb == "FF969696"
    # Fuente y formato propios de los datos; el formato que exige una regla prevalece.
    assert (sheet["C2"].font.name, sheet["C2"].font.sz, sheet["C2"].number_format) == ("Arial", 9, "0")
    assert sheet["A2"].number_format == '"$" #,##0.00'
    assert sheet["B2"].number_format == "#,##0.00"
    # Anchos originales; la columna nueva se ajusta a su contenido.
    assert sheet.column_dimensions["A"].width == 29
    assert sheet.column_dimensions["B"].width == 18


def test_without_original_appearance_headers_are_bold():
    sheet = load_workbook(BytesIO(build_output(pd.DataFrame({"A": [1]}), []))).active
    assert sheet.title == "Datos" and sheet["A1"].font.b
