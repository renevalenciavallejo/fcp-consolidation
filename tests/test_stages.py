"""Pruebas históricas: cada etapa contra resultados manuales reales.

Estructura de cada caso (los datos reales no se suben a git):

    tests/cases/stage_NN/<semana>/
        inputs/<id_entrada>/*.xlsx   # una carpeta por entrada de la etapa
        expected.xlsx                # resultado manual real

En una entrada por fecha, cada archivo se nombra con su fecha: aaaa-mm-dd.xlsx.

Una diferencia no es automáticamente un error de la herramienta: puede ser un
error del proceso manual. Se analiza y se documenta.

Las celdas con fórmulas del archivo generado no se comparan: openpyxl no las
calcula y el archivo aún no ha pasado por Excel, así que no tienen valor.
"""

from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.runner import run_stage
from core.stage import input_specs, stage_prefix
from stages import STAGES

CASES_DIR = Path(__file__).parent / "cases"

# Diferencia máxima aceptada entre dos números (por redondeos de Excel).
NUMBER_TOLERANCE = 0.005

# Cantidad máxima de diferencias que se muestran en el reporte.
MAX_DIFFERENCES_SHOWN = 50


def find_cases() -> list:
    cases = []
    for stage in STAGES:
        folder = CASES_DIR / f"stage_{stage_prefix(stage)[1:3]}"
        if not folder.is_dir():
            continue
        for case in sorted(p for p in folder.iterdir() if p.is_dir()):
            cases.append(pytest.param(stage, case, id=f"{folder.name}/{case.name}"))
    return cases


def read_rows(source, keep_formulas: bool = False) -> list[list]:
    sheet = load_workbook(source, data_only=not keep_formulas).worksheets[0]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def is_formula(value) -> bool:
    return isinstance(value, str) and value.startswith("=")


def compare_rows(obtained: list[list], expected: list[list]) -> list[str]:
    """Compara celda por celda. Devuelve una descripción por diferencia (fila, columna, valores)."""
    differences = []
    if obtained[:1] != expected[:1]:
        differences.append(f"Encabezados distintos: obtenido {obtained[:1]}, esperado {expected[:1]}")
        return differences
    header = expected[0] if expected else []
    if len(obtained) != len(expected):
        differences.append(f"Cantidad de filas: obtenido {len(obtained) - 1}, esperado {len(expected) - 1}")
    for row_number, (got_row, exp_row) in enumerate(zip(obtained[1:], expected[1:]), start=2):
        for column, got, exp in zip(header, got_row, exp_row):
            if not is_formula(got) and not same_value(got, exp):
                differences.append(f"Fila {row_number}, columna '{column}': obtenido {got!r}, esperado {exp!r}")
    return differences


def same_value(got, exp) -> bool:
    if got in (None, "") and exp in (None, ""):
        return True
    if isinstance(got, (int, float)) and isinstance(exp, (int, float)):
        return abs(got - exp) <= NUMBER_TOLERANCE
    if isinstance(got, (datetime, date)) and isinstance(exp, (datetime, date)):
        return got == exp
    return str(got).strip() == str(exp).strip()


@pytest.mark.parametrize("stage, case", find_cases())
def test_historical_case(stage, case: Path):
    files = {
        folder.name: sorted(folder.glob("*.xlsx"))
        for folder in (case / "inputs").iterdir()
        if folder.is_dir()
    }
    for spec in input_specs(stage):
        if spec.per_date and spec.key in files:
            files[spec.key] = {date.fromisoformat(path.stem): path for path in files[spec.key]}
    result = run_stage(stage, files)
    obtained = read_rows(BytesIO(result.output_bytes), keep_formulas=True)
    differences = compare_rows(obtained, read_rows(case / "expected.xlsx"))
    shown = "\n".join(differences[:MAX_DIFFERENCES_SHOWN])
    assert not differences, f"{len(differences)} diferencias en {case}:\n{shown}"
