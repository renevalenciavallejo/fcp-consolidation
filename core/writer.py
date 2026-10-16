"""Construcción de archivos de Excel con el formato estándar de la herramienta."""

from copy import copy
from datetime import date, datetime
from io import BytesIO
from typing import Any

import re
from dataclasses import dataclass, field

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.audit import Finding, data_columns, is_empty
from core.settings import (
    DATE_FORMAT_EXCEL,
    MAX_COLUMN_WIDTH,
    NUMBER_FORMAT_EXCEL,
    SEVERITY_FILLS,
    SOURCE_FILE,
    SOURCE_ROW,
)

# Referencia a otra columna dentro de una fórmula: {Nombre de la columna}.
COLUMN_REFERENCE = re.compile(r"\{([^}]+)\}")

# Ancho (en caracteres) de las columnas con fórmulas, cuyo texto no es lo que se ve.
FORMULA_COLUMN_WIDTH = 16


@dataclass(frozen=True)
class ExcelFormula:
    """Fórmula de Excel que se escribe en cada fila del archivo de salida.

    Las columnas se nombran entre llaves y se convierten en la celda de la misma
    fila al escribir, por ejemplo "={VR. BRUTO}-{Importe}" -> "=E2-J2". Así la
    fórmula sigue siendo correcta aunque cambie la posición de las columnas.
    """

    template: str
    number_format: str | None = None


# Orden de prioridad de las severidades cuando una celda tiene varios hallazgos.
SEVERITY_PRIORITY = list(SEVERITY_FILLS)

# Formato de número de Excel que equivale a «sin formato».
GENERAL_FORMAT = "General"


@dataclass(frozen=True)
class CellStyle:
    """Apariencia de una celda: fuente, relleno, bordes, alineación y formato de número."""

    font: Font
    fill: PatternFill
    border: Border
    alignment: Alignment
    number_format: str


@dataclass
class SheetStyle:
    """Apariencia de la hoja original, para repetirla en el archivo de salida.

    Así el usuario ve el archivo procesado igual que el que cargó y se fija solo
    en los datos. Los estilos se guardan por nombre de columna, porque una etapa
    puede agregar columnas y correr las demás de posición.
    """

    sheet_name: str
    header: dict[str, CellStyle] = field(default_factory=dict)   # Celda del encabezado
    data: dict[str, CellStyle] = field(default_factory=dict)     # Primera celda de datos
    widths: dict[str, float] = field(default_factory=dict)       # Ancho de la columna


def capture_style(sheet: Worksheet, header_row: int, positions: dict[str, int]) -> SheetStyle:
    """Toma la apariencia de una hoja: nombre, encabezados, primera fila de datos y anchos.

    `positions` indica la posición (desde 0) de cada columna en la hoja.
    """
    widths = _column_widths(sheet)
    style = SheetStyle(sheet.title)
    for column, position in positions.items():
        style.header[column] = _cell_style(sheet.cell(header_row, position + 1))
        style.data[column] = _cell_style(sheet.cell(header_row + 1, position + 1))
        if position + 1 in widths:
            style.widths[column] = widths[position + 1]
    return style


def build_output(
    df: pd.DataFrame,
    findings: list[Finding],
    column_formats: dict[str, str] | None = None,
    style: SheetStyle | None = None,
) -> bytes:
    """Genera el archivo de salida: datos procesados sin columnas internas.

    `column_formats` indica un formato de Excel por columna (por ejemplo, moneda);
    las demás columnas usan el formato según el tipo de valor.
    `style` es la apariencia del archivo original (nombre de la hoja, encabezados,
    fuentes y anchos), que se repite en la salida.
    Las celdas con hallazgos se resaltan según la severidad.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = style.sheet_name if style else "Datos"
    write_table(sheet, df[data_columns(df)], column_formats, style)
    _highlight_findings(sheet, df, findings)
    return to_bytes(workbook)


def write_table(
    sheet: Worksheet,
    df: pd.DataFrame,
    column_formats: dict[str, str] | None = None,
    style: SheetStyle | None = None,
) -> None:
    """Escribe un DataFrame con filtros, paneles inmovilizados y anchos ajustados.

    Sin `style`, los encabezados van en negrita. Con `style`, cada columna toma la
    apariencia que tenía en el archivo original; una columna nueva toma la de la
    columna a su izquierda.
    """
    columns = list(df.columns)
    sheet.append([str(c) for c in columns])
    header_styles, data_styles = _styles_by_column(columns, style)
    for cell, column in zip(sheet[1], columns):
        if header_styles[column]:
            _apply_style(cell, header_styles[column])
        elif style is None:
            cell.font = Font(bold=True)

    letters = {column: get_column_letter(i + 1) for i, column in enumerate(columns)}
    for row_number, values in enumerate(df.itertuples(index=False), start=2):
        sheet.append([_excel_value(v, letters, row_number) for v in values])

    formats = column_formats or {}
    formula_formats = _formula_formats(df)
    for row in sheet.iter_rows(min_row=2):
        for cell, column in zip(row, columns):
            original = data_styles[column]
            if original:
                _apply_style(cell, original)
            # Los formatos que exige una regla (moneda, fórmula, fecha) prevalecen sobre el original.
            if column in formula_formats:
                if formula_formats[column]:
                    cell.number_format = formula_formats[column]
                continue
            is_number = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            if column in formats and is_number:
                cell.number_format = formats[column]
            elif isinstance(cell.value, (datetime, date)):
                cell.number_format = DATE_FORMAT_EXCEL
            elif style and column in style.data and original.number_format != GENERAL_FORMAT:
                # Formato propio de la columna en el original (no el heredado de su vecina).
                cell.number_format = original.number_format
            elif isinstance(cell.value, float):
                cell.number_format = NUMBER_FORMAT_EXCEL

    if columns:
        sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    _adjust_widths(sheet)
    if style:
        for i, column in enumerate(columns):
            if column in style.widths:
                sheet.column_dimensions[get_column_letter(i + 1)].width = style.widths[column]


def to_bytes(workbook: Workbook) -> bytes:
    """Guarda el libro en memoria y devuelve su contenido."""
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _excel_value(value: Any, letters: dict[str, str] | None = None, row_number: int = 0) -> Any:
    """Prepara un valor para Excel: vacíos como celda vacía, fechas de pandas como datetime
    y fórmulas con las referencias a columnas convertidas en celdas de la fila."""
    if isinstance(value, ExcelFormula):
        return _resolve_formula(value, letters or {}, row_number)
    if is_empty(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        return value.item()  # Tipos de numpy -> tipos de Python
    return value


def _resolve_formula(formula: ExcelFormula, letters: dict[str, str], row_number: int) -> str:
    def cell_reference(match: re.Match) -> str:
        column = match.group(1)
        if column not in letters:
            raise ValueError(f"La fórmula '{formula.template}' usa la columna '{column}', que no existe.")
        return f"{letters[column]}{row_number}"

    return COLUMN_REFERENCE.sub(cell_reference, formula.template)


def _formula_formats(df: pd.DataFrame) -> dict[str, str | None]:
    """Columnas que contienen fórmulas y el formato de número que llevan."""
    formats = {}
    for column in df.columns:
        first = next((v for v in df[column] if isinstance(v, ExcelFormula)), None)
        if first is not None:
            formats[column] = first.number_format
    return formats


def _cell_style(cell: Cell) -> CellStyle:
    return CellStyle(copy(cell.font), copy(cell.fill), copy(cell.border), copy(cell.alignment), cell.number_format)


def _apply_style(cell: Cell, style: CellStyle) -> None:
    """Aplica fuente, relleno, bordes y alineación (el formato de número se decide aparte)."""
    cell.font = style.font
    cell.fill = style.fill
    cell.border = style.border
    cell.alignment = style.alignment


def _styles_by_column(
    columns: list[str], style: SheetStyle | None
) -> tuple[dict[str, CellStyle | None], dict[str, CellStyle | None]]:
    """Estilo de encabezado y de datos de cada columna de salida.

    Una columna que no estaba en el original (por ejemplo, una calculada) toma el
    estilo de la columna a su izquierda, para que no se vea distinta de las demás.
    """
    header: dict[str, CellStyle | None] = {}
    data: dict[str, CellStyle | None] = {}
    previous = None
    for column in columns:
        if style and column in style.header:
            previous = column
        header[column] = style.header.get(previous) if style and previous else None
        data[column] = style.data.get(previous) if style and previous else None
    return header, data


def _column_widths(sheet: Worksheet) -> dict[int, float]:
    """Ancho de cada columna con ancho definido. Excel guarda rangos de columnas
    con el mismo ancho (por ejemplo, E a J), así que se reparten una por una."""
    widths = {}
    for letter, dimension in sheet.column_dimensions.items():
        if not dimension.customWidth or not dimension.width:
            continue
        first = dimension.min or column_index_from_string(letter)
        for number in range(first, (dimension.max or first) + 1):
            widths[number] = dimension.width
    return widths


def _adjust_widths(sheet: Worksheet) -> None:
    for column_cells in sheet.columns:
        longest = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            if isinstance(cell.value, (datetime, date)):
                length = 10
            elif isinstance(cell.value, float):
                length = len(f"$ {cell.value:,.2f}")
            elif isinstance(cell.value, str) and cell.value.startswith("="):
                length = FORMULA_COLUMN_WIDTH
            else:
                length = len(str(cell.value))
            longest = max(longest, length)
        letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[letter].width = min(longest + 2, MAX_COLUMN_WIDTH)


def _highlight_findings(sheet: Worksheet, df: pd.DataFrame, findings: list[Finding]) -> None:
    """Colorea las celdas que tienen hallazgos. Si el registro o la columna ya no están, se omite."""
    if SOURCE_FILE not in df.columns or SOURCE_ROW not in df.columns:
        return
    output_rows = {
        (file, int(row)): position + 2  # +2: encabezado y Excel empieza en 1
        for position, (file, row) in enumerate(zip(df[SOURCE_FILE], df[SOURCE_ROW]))
    }
    output_columns = {c: i + 1 for i, c in enumerate(data_columns(df))}

    severity_by_cell: dict[tuple[int, int], str] = {}
    for finding in findings:
        if finding.severity not in SEVERITY_FILLS:
            continue
        row = output_rows.get((finding.source_file, finding.source_row))
        column = output_columns.get(finding.column)
        if row is None or column is None:
            continue
        current = severity_by_cell.get((row, column))
        if current is None or SEVERITY_PRIORITY.index(finding.severity) < SEVERITY_PRIORITY.index(current):
            severity_by_cell[(row, column)] = finding.severity

    for (row, column), severity in severity_by_cell.items():
        color = SEVERITY_FILLS[severity]
        sheet.cell(row=row, column=column).fill = PatternFill("solid", start_color=color, end_color=color)
