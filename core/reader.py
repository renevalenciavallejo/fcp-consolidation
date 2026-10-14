"""Lectura de archivos de Excel y validación estricta de su estructura.

Las celdas se leen con openpyxl tal como están guardadas en Excel (texto,
número o fecha), sin dejar que pandas adivine tipos. Así no se pierden ceros a
la izquierda ni se confunden formatos; la conversión explícita la hacen las
funciones de core/cleaning.py.
"""

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.audit import is_empty
from core.cleaning import clean_text
from core.settings import SOURCE_FILE, SOURCE_ROW
from core.stage import InputSpec


# Cantidad de filas iniciales en las que se buscan encabezados corridos de lugar.
HEADER_SEARCH_ROWS = 20


@dataclass
class ReadResult:
    """Resultado de leer un archivo: los datos (si la estructura es válida) y los problemas."""

    file_name: str
    data: pd.DataFrame | None
    problems: list[str]
    # Filas anteriores a los encabezados (títulos del reporte): (fila en Excel, contenido por columna).
    # No son registros; el ejecutor las registra como eliminadas para que el log las muestre.
    title_rows: list[tuple[int, dict]] = field(default_factory=list)

    @property
    def rows(self) -> int:
        return 0 if self.data is None else len(self.data)


def file_name(file: Any) -> str:
    """Nombre del archivo, ya sea una ruta, un archivo cargado en Streamlit o un BytesIO."""
    if isinstance(file, (str, Path)):
        return Path(file).name
    return Path(getattr(file, "name", "archivo.xlsx")).name


def read_file(file: Any, spec: InputSpec) -> ReadResult:
    """Lee un archivo de una entrada y valida su estructura.

    Si hay algún problema de estructura, `data` es None y `problems` explica qué
    falta o sobra. El archivo original solo se lee, nunca se modifica.
    """
    name = file_name(file)
    try:
        workbook = load_workbook(_as_source(file), data_only=True)
    except Exception:
        return ReadResult(name, None, [f"[EST-001] {name}: no se pudo abrir. Verifique que sea un archivo .xlsx válido."])

    if spec.sheet is None:
        # Sin nombre configurado se usa la primera hoja: el nombre puede variar entre exportaciones.
        sheet_name = workbook.sheetnames[0]
    elif spec.sheet in workbook.sheetnames:
        sheet_name = spec.sheet
    else:
        sheets = ", ".join(workbook.sheetnames)
        return ReadResult(
            name, None, [f"[EST-002] {name}: no tiene la hoja '{spec.sheet}'. Hojas encontradas: {sheets}."]
        )

    rows = list(workbook[sheet_name].iter_rows(min_row=1, min_col=1, values_only=True))
    workbook.close()

    header = rows[spec.header_row - 1] if len(rows) >= spec.header_row else ()
    data_rows = _trim_trailing_empty(rows[spec.header_row:])
    width = max([len(header), *(len(r) for r in data_rows)], default=0)

    problems: list[str] = []
    positions: dict[str, int] = {}
    for position in range(width):
        title = header[position] if position < len(header) else None
        title = clean_text(str(title)) if title is not None else ""
        if title == "":
            if any(position < len(r) and not _empty_cell(r[position]) for r in data_rows):
                letter = get_column_letter(position + 1)
                problems.append(f"[EST-006] {name}: la columna {letter} tiene datos pero no tiene encabezado.")
            continue
        if title in positions:
            problems.append(f"[EST-005] {name}: el encabezado '{title}' está repetido.")
            continue
        positions[title] = position

    missing = [c for c in spec.columns if c not in positions]
    found_row = _find_header_row(rows, spec) if missing else None
    if found_row is not None:
        # Si los encabezados están en otra fila, los demás problemas serían ruido.
        return ReadResult(name, None, [
            f"[EST-003] {name}: los encabezados están en la fila {found_row} de la hoja '{sheet_name}', "
            f"pero se esperaban en la fila {spec.header_row}. Verifique que no se hayan agregado "
            "o eliminado filas al inicio del archivo."
        ])
    if missing:
        problems.append(
            f"[EST-003] {name}: faltan las columnas {_quote(missing)} "
            f"en la fila {spec.header_row} de la hoja '{sheet_name}'."
        )
    extra = [c for c in positions if c not in spec.columns]
    if extra and not spec.allow_extra_columns:
        problems.append(f"[EST-004] {name}: tiene columnas que no se esperaban: {_quote(extra)}.")

    if not problems and spec.check_column_order:
        problems += _check_order(name, list(positions), list(spec.columns))

    if problems:
        return ReadResult(name, None, problems)

    columns = list(positions)
    records = []
    for offset, row in enumerate(data_rows):
        record = [row[positions[c]] if positions[c] < len(row) else None for c in columns]
        # Fila real en Excel: la que el usuario puede buscar en su archivo.
        records.append([*record, name, spec.header_row + 1 + offset])
    data = pd.DataFrame(records, columns=[*columns, SOURCE_FILE, SOURCE_ROW], dtype=object)
    return ReadResult(name, data, [], _title_rows(rows[: spec.header_row - 1], header))


def combine(results: list[ReadResult], spec: InputSpec) -> pd.DataFrame:
    """Une en un solo DataFrame los archivos de una entrada que tienen la misma estructura."""
    frames = [r.data for r in results if r.data is not None]
    if not frames:
        return pd.DataFrame(columns=[*spec.columns, SOURCE_FILE, SOURCE_ROW], dtype=object)
    return pd.concat(frames, ignore_index=True)


def _title_rows(rows: list[tuple], header: tuple) -> list[tuple[int, dict]]:
    """Contenido de las filas de título, en la columna de Excel donde está cada valor.

    Se usa el nombre del encabezado de esa columna para que en la hoja Eliminados
    cada valor aparezca bajo la misma columna que tiene en el archivo original.
    """
    result = []
    for number, row in enumerate(rows, start=1):
        content = {}
        for position, value in enumerate(row):
            title = header[position] if position < len(header) else None
            column = clean_text(str(title)) if title is not None else ""
            if column == "":
                if _empty_cell(value):
                    continue
                column = f"Columna {get_column_letter(position + 1)}"
            content[column] = value
        result.append((number, content))
    return result


def _check_order(name: str, found: list[str], expected: list[str]) -> list[str]:
    """Compara el orden de las columnas esperadas con el del archivo. Regla EST-008.

    Se informa la primera posición distinta, que es la que el usuario debe revisar.
    """
    found = [c for c in found if c in expected]  # Las columnas adicionales permitidas no cuentan
    for position, (got, wanted) in enumerate(zip(found, expected), start=1):
        if got != wanted:
            return [
                f"[EST-008] {name}: las columnas no están en el orden esperado. "
                f"En la posición {position} se esperaba '{wanted}' y se encontró '{got}'."
            ]
    return []


def _find_header_row(rows: list[tuple], spec: InputSpec) -> int | None:
    """Busca en las primeras filas una que tenga la mayoría de los encabezados esperados.

    Sirve para explicar al usuario un error común: filas de título agregadas o
    borradas al inicio, que corren los encabezados de lugar.
    """
    expected = set(spec.columns)
    best_row, best_count = None, len(expected) // 2
    for number, row in enumerate(rows[:HEADER_SEARCH_ROWS], start=1):
        names = {clean_text(str(v)) for v in row if v is not None}
        count = len(expected & names)
        if count > best_count:
            best_row, best_count = number, count
    return best_row if best_row != spec.header_row else None


def _as_source(file: Any) -> Any:
    """Devuelve algo que openpyxl pueda abrir, empezando siempre desde el inicio del archivo."""
    if isinstance(file, (str, Path)):
        return Path(file)
    if hasattr(file, "seek"):
        file.seek(0)
    if hasattr(file, "getvalue"):
        return BytesIO(file.getvalue())
    return file


def _empty_cell(value: Any) -> bool:
    return is_empty(value) or (isinstance(value, str) and clean_text(value) == "")


def _trim_trailing_empty(rows: list[tuple]) -> list[tuple]:
    """Quita las filas vacías del final de la hoja.

    Excel a veces reporta filas con formato pero sin datos después del último
    registro; no son registros y no deben aparecer como eliminados.
    """
    end = len(rows)
    while end > 0 and all(_empty_cell(v) for v in rows[end - 1]):
        end -= 1
    return rows[:end]


def _quote(names: list[str]) -> str:
    return ", ".join(f"'{n}'" for n in names)
